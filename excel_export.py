from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter
PURPLE='72007E';LIGHT='F4ECF6';BLUE='17598A'
def build_excel(payload,pre,ctl,rows,roi_events,pp_events):
    wb=Workbook();ws=wb.active;ws.title='EMI Schedule';ws.sheet_view.showGridLines=False
    ws.merge_cells('A1:I1');ws['A1']='ASBS & CO - Loan Repayment Schedule';ws['A1'].font=Font(size=17,bold=True,color='FFFFFF');ws['A1'].fill=PatternFill('solid',fgColor=PURPLE);ws['A1'].alignment=Alignment(horizontal='center');ws.row_dimensions[1].height=28
    meta=[('Borrower',payload['borrower']),('Reference',payload['ref']),('Principal',payload['principal']),('Initial EMI',ctl['initial_emi']),('Pre-EMI',pre),('Total Interest',ctl['total_interest']),('Part Prepayments',ctl['total_prepayment']),('Payoff Date',ctl['payoff_date']),('Control Status',ctl['status']),('Principal Variance',ctl['principal_variance'])]
    for i,(k,v) in enumerate(meta,3):ws.cell(i,1,k).font=Font(bold=True,color=PURPLE);ws.cell(i,2,v)
    hdr=14;headers=['EMI #','Due Date','Opening Principal','ROI (% p.a.)','Interest Component','Principal Component','EMI','Part Prepayment','Closing Principal']
    for j,h in enumerate(headers,1):c=ws.cell(hdr,j,h);c.fill=PatternFill('solid',fgColor=PURPLE);c.font=Font(bold=True,color='FFFFFF');c.alignment=Alignment(horizontal='center')
    for i,r in enumerate(rows,hdr+1):
        vals=[r.emi_no,r.due_date,r.opening_principal,r.roi,r.interest_component,r.principal_component,r.emi,r.part_prepayment,r.closing_principal]
        for j,v in enumerate(vals,1):ws.cell(i,j,v)
        ws.cell(i,2).number_format='dd-mm-yyyy';ws.cell(i,4).number_format='0.000%'
        for j in [3,5,6,7,8,9]:ws.cell(i,j).number_format='₹#,##0.00;[Red](₹#,##0.00);-'
    ws.freeze_panes='A15';ws.auto_filter.ref=f'A14:I{14+len(rows)}'
    widths=[10,15,20,15,20,20,16,18,20]
    for i,w in enumerate(widths,1):ws.column_dimensions[get_column_letter(i)].width=w
    for title,events,heads in [('ROI Changes',roi_events,['Effective Date','New ROI (% p.a.)','Reference / Note']),('Part Prepayments',pp_events,['Date','Amount','Effect','Reference / Note'])]:
        sh=wb.create_sheet(title);sh.sheet_view.showGridLines=False
        for j,h in enumerate(heads,1):c=sh.cell(1,j,h);c.fill=PatternFill('solid',fgColor=PURPLE);c.font=Font(bold=True,color='FFFFFF')
        for i,e in enumerate(events,2):
            vals=[e.get('date'),e.get('rate'),e.get('note')] if title=='ROI Changes' else [e.get('date'),e.get('amount'),e.get('effect'),e.get('note')]
            for j,v in enumerate(vals,1):sh.cell(i,j,v)
        for j in range(1,len(heads)+1):sh.column_dimensions[get_column_letter(j)].width=24
    out=BytesIO();wb.save(out);return out.getvalue()
