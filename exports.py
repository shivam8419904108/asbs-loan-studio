from io import BytesIO
from pathlib import Path
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4,landscape
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate,Table,TableStyle,Paragraph,Spacer,Image
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
PURPLE='72007E'; MONEY='₹#,##0.00;[Red](₹#,##0.00);-'
def build_pdf(logo,payload,pre,ctl,rows):
    out=BytesIO();doc=SimpleDocTemplate(out,pagesize=landscape(A4),leftMargin=8*mm,rightMargin=8*mm,topMargin=8*mm,bottomMargin=8*mm);s=getSampleStyleSheet();s.add(ParagraphStyle(name='T',parent=s['Title'],textColor=colors.HexColor('#'+PURPLE),fontSize=17));story=[]
    if Path(logo).exists():story.append(Image(str(logo),width=50*mm,height=15*mm))
    story += [Paragraph('Loan Repayment Schedule',s['T']),Spacer(1,3*mm)]
    meta=[['Name',payload.get('name') or '-','Account Number',payload.get('account_number') or '-'],['Principal',f"{payload['principal']:,.2f}",'Initial EMI',f"{ctl['initial_emi']:,.2f}"],['Pre-EMI',f'{pre:,.2f}','Total Interest',f"{ctl['total_interest']:,.2f}"],['Part Prepayments',f"{ctl['total_prepayment']:,.2f}",'Control',ctl['status']]]
    t=Table(meta,colWidths=[32*mm,53*mm,35*mm,55*mm]);t.setStyle(TableStyle([('BACKGROUND',(0,0),(0,-1),colors.HexColor('#F2E8F4')),('BACKGROUND',(2,0),(2,-1),colors.HexColor('#F2E8F4')),('GRID',(0,0),(-1,-1),.25,colors.lightgrey),('PADDING',(0,0),(-1,-1),5)]));story += [t,Spacer(1,4*mm)]
    data=[['EMI #','Due Date','ROI (% p.a.)','Interest Component','Principal Component','EMI','Closing Principal']]+[[r.emi_no,r.due_date.strftime('%d-%m-%Y'),f'{r.roi*100:.2f}%',f'{r.interest_component:,.2f}',f'{r.principal_component:,.2f}',f'{r.emi:,.2f}',f'{r.closing_principal:,.2f}'] for r in rows]
    t=Table(data,repeatRows=1,colWidths=[15*mm,27*mm,27*mm,36*mm,38*mm,30*mm,38*mm]);t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#'+PURPLE)),('TEXTCOLOR',(0,0),(-1,0),colors.white),('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#FAF6FB')]),('GRID',(0,0),(-1,-1),.25,colors.lightgrey),('ALIGN',(2,1),(-1,-1),'RIGHT'),('FONTSIZE',(0,0),(-1,-1),8)]));story.append(t);doc.build(story);return out.getvalue()
def build_excel(payload,pre,ctl,rows,roi_events,pp_events):
    wb=Workbook();ws=wb.active;ws.title='EMI Schedule';ws.sheet_view.showGridLines=False;ws.merge_cells('A1:I1');ws['A1']='ASBS & CO - Loan Repayment Schedule';ws['A1'].font=Font(size=17,bold=True,color='FFFFFF');ws['A1'].fill=PatternFill('solid',fgColor=PURPLE);ws['A1'].alignment=Alignment(horizontal='center')
    meta=[('Name',payload.get('name','')),('Account Number',payload.get('account_number','')),('Principal',payload['principal']),('Entered Tenure',payload['tenure']),('Schedule Rows',ctl['schedule_rows']),('Initial EMI',ctl['initial_emi']),('Pre-EMI',pre),('Total Interest',ctl['total_interest']),('Part Prepayments',ctl['total_prepayment']),('Control Status',ctl['status']),('Annual Basis','365 days only')]
    for row,(label,value) in enumerate(meta,3):ws.cell(row,1,label).font=Font(bold=True,color=PURPLE);ws.cell(row,2,value)
    hr=15;heads=['EMI #','Due Date','Opening Principal','ROI (% p.a.)','Interest Component','Principal Component','EMI','Part Prepayment','Closing Principal']
    for c,h in enumerate(heads,1):cell=ws.cell(hr,c,h);cell.fill=PatternFill('solid',fgColor=PURPLE);cell.font=Font(bold=True,color='FFFFFF')
    for row,item in enumerate(rows,hr+1):
        vals=[item.emi_no,item.due_date,item.opening_principal,item.roi,item.interest_component,item.principal_component,item.emi,item.part_prepayment,item.closing_principal]
        for c,v in enumerate(vals,1):ws.cell(row,c,v)
        ws.cell(row,2).number_format='dd-mm-yyyy';ws.cell(row,4).number_format='0.00%'
        for c in [3,5,6,7,8,9]:ws.cell(row,c).number_format=MONEY
    ws.freeze_panes='A16';ws.auto_filter.ref=f'A15:I{hr+len(rows)}'
    for c,w in enumerate([10,15,20,15,20,20,16,18,20],1):ws.column_dimensions[get_column_letter(c)].width=w
    for title,events,heads in [('ROI Changes',roi_events,['Effective Date','New ROI (% p.a.)','Note']),('Part Prepayments',pp_events,['Date','Amount','Effect','Note'])]:
        sh=wb.create_sheet(title)
        for c,h in enumerate(heads,1):cell=sh.cell(1,c,h);cell.fill=PatternFill('solid',fgColor=PURPLE);cell.font=Font(bold=True,color='FFFFFF')
        for row,e in enumerate(events,2):
            vals=[e.get('date'),e.get('rate'),e.get('note')] if title=='ROI Changes' else [e.get('date'),e.get('amount'),e.get('effect'),e.get('note')]
            for c,v in enumerate(vals,1):sh.cell(row,c,v)
    out=BytesIO();wb.save(out);return out.getvalue()
